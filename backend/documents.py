"""Turn the Markdown Ava writes into downloadable .docx and .pdf files.

The PDF uses a LaTeX-résumé style: a centered name, a compact contact line,
and serif section headings in caps underlined by a full-width rule.
"""

import os
import re
import unicodedata
import uuid

from docx import Document as Docx
from fpdf import FPDF
from fpdf.enums import XPos, YPos

GEN_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generated")
os.makedirs(GEN_DIR, exist_ok=True)

_BOLD = re.compile(r"\*\*(.+?)\*\*")

# PDF core fonts are latin-1 only; fold typographic characters to ASCII.
_PDF_SUBS = {"•": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "…": "..."}


def _safe(text: str) -> str:
    out = []
    for ch in text:
        if ch in _PDF_SUBS:
            out.append(_PDF_SUBS[ch])
            continue
        cat = unicodedata.category(ch)
        if cat == "Zs":        # any Unicode space (incl. narrow no-break) -> normal space
            out.append(" ")
        elif cat == "Pd":      # any dash (en/em/figure/…) -> hyphen
            out.append("-")
        else:
            out.append(ch)
    return "".join(out).encode("latin-1", "replace").decode("latin-1")


def _add_runs(paragraph, text: str) -> None:
    """Add text to a docx paragraph, turning **...** into bold runs."""
    pos = 0
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _markdown_to_docx(title: str, markdown_text: str, out_path: str) -> None:
    doc = Docx()
    if title:
        doc.add_heading(title, level=0)
    for raw in (markdown_text or "").replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", line))
        else:
            _add_runs(doc.add_paragraph(), line)
    doc.save(out_path)


# ---------------------------------------------------------------------------
# PDF (LaTeX-résumé style)
# ---------------------------------------------------------------------------

def _line(pdf, text, height, markdown=False, align="L"):
    pdf.multi_cell(0, height, text, markdown=markdown, align=align,
                   new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _rule(pdf):
    y = pdf.get_y() + 0.4
    pdf.set_draw_color(120, 120, 120)
    pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
    pdf.ln(2.4)


def _bullet(pdf, text):
    left = pdf.l_margin
    y = pdf.get_y()
    pdf.set_fill_color(70, 70, 70)
    pdf.ellipse(left + 1.8, y + 1.9, 1.5, 1.5, style="F")  # small drawn disc
    pdf.set_left_margin(left + 6)
    pdf.set_x(left + 6)
    pdf.multi_cell(0, 5.6, text, markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_left_margin(left)
    pdf.set_x(left)


def _split_header(lines, name_fallback):
    """Pull the name (first '# ') and the following contact line out of the body."""
    name, contact, start = None, None, 0
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("# ") and not s.startswith("## "):
            name = s[2:].strip().replace("**", "")
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and not lines[i].strip().startswith("#"):
                contact = re.sub(r"^\s*[-*]\s+", "", lines[i].strip()).replace("**", "")
                i += 1
            start = i
            break
        i += 1
    if not name:
        name = re.sub(r"\s+(CV|Cover letter)$", "", name_fallback or "", flags=re.I).strip()
        start = 0
    return name, contact, start


def _markdown_to_pdf(name_fallback: str, markdown_text: str, out_path: str) -> None:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(18, 15, 18)
    pdf.add_page()

    lines = [l.rstrip() for l in (markdown_text or "").replace("\r\n", "\n").split("\n")]
    name, contact, start = _split_header(lines, name_fallback)

    pdf.set_font("Times", "B", 22)
    _line(pdf, _safe(name), 9, align="C")
    if contact:
        pdf.set_font("Times", "", 10)
        pdf.set_text_color(90, 90, 90)
        _line(pdf, _safe(contact), 5.5, align="C")
        pdf.set_text_color(0, 0, 0)
    pdf.ln(2.5)

    for raw in lines[start:]:
        s = raw.strip()
        if not s:
            pdf.ln(1.8)
            continue
        if s.startswith("## ") or (s.startswith("# ") and not s.startswith("## ")):
            heading = s.lstrip("#").strip()
            pdf.ln(1.5)
            pdf.set_font("Times", "B", 12)
            _line(pdf, _safe(heading.upper()), 6)
            _rule(pdf)
        elif s.startswith("### "):
            pdf.set_font("Times", "B", 11)
            _line(pdf, _safe(s[4:].strip()), 5.8, markdown=True)
        elif re.match(r"^\s*[-*]\s+", s):
            pdf.set_font("Times", "", 11)
            _bullet(pdf, _safe(re.sub(r"^\s*[-*]\s+", "", s)))
        else:
            pdf.set_font("Times", "", 11)
            _line(pdf, _safe(s), 5.8, markdown=True)

    pdf.output(out_path)


def build_document(kind: str, title: str, content_markdown: str) -> tuple[str, str, str]:
    """Write the document to disk as .docx and .pdf. Returns (id, docx_path, pdf_path)."""
    doc_id = uuid.uuid4().hex
    docx_path = os.path.join(GEN_DIR, f"{doc_id}.docx")
    pdf_path = os.path.join(GEN_DIR, f"{doc_id}.pdf")
    heading = title or ("CV" if kind == "cv" else "Cover letter")
    _markdown_to_docx(heading, content_markdown, docx_path)
    _markdown_to_pdf(heading, content_markdown, pdf_path)
    return doc_id, docx_path, pdf_path
